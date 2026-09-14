#!/usr/bin/env python3

"""
Cut football clips from one source video and extract their audio.

Project structure:

footballBrains/
├── setup/
│   ├── videoCutter.py
│   ├── videoSpreadsheet.xlsx
│   └── footballGoalsDVD.mp4
├── video/
└── audio/
    └── raw/

The script:
1. Reads videoSpreadsheet.xlsx.
2. Cuts each requested MP4 clip.
3. Saves clips to ../video/.
4. Extracts each clip's audio as a WAV file.
5. Saves WAV files to ../audio/raw/.
"""

import os
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


# ----- Project paths -----

setupDir = Path(__file__).resolve().parent
projectDir = setupDir.parent

spreadsheetPath = setupDir / "videoSpreadsheet.xlsx"
sourceVideoName = "footballGoalsDVD.mp4"
sourceVideoPath = setupDir / sourceVideoName

videoDir = projectDir / "video"
audioDir = projectDir / "audio" / "raw"

videoDir.mkdir(parents=True, exist_ok=True)
audioDir.mkdir(parents=True, exist_ok=True)


# ----- Video encoding settings -----

videoCodec = "libx264"
audioCodec = "aac"
audioBitrate = "192k"
pixelFormat = "yuv420p"


# ----- Audio extraction settings -----

audioSampleRate = "48000"
audioChannels = "2"
audioCodecPcm = "pcm_s16le"


# ----- Helper functions -----


def checkBinary(name):
    """Check that a command-line program is installed."""

    binaryPath = shutil.which(name)

    if binaryPath is None:
        raise SystemExit(
            f"Missing dependency: {name}\n"
            "Install FFmpeg, then try again."
        )

    return binaryPath


def runCommand(command):
    """Print and run a command."""

    print(
        "RUN:",
        " ".join(map(str, command)),
        flush=True
    )

    subprocess.run(
        command,
        check=True
    )


def parseTimeToSeconds(startTime):
    """
    Convert hh:mm:ss, mm:ss, or seconds to seconds.
    """

    text = str(startTime).strip()

    if not text:
        raise ValueError(
            "startTime cannot be blank"
        )

    try:
        parts = [
            float(part.strip())
            for part in text.split(":")
        ]

    except ValueError as error:
        raise ValueError(
            f"Invalid startTime: {startTime!r}"
        ) from error

    if len(parts) == 3:
        hours, minutes, seconds = parts

    elif len(parts) == 2:
        hours = 0.0
        minutes, seconds = parts

    elif len(parts) == 1:
        hours = 0.0
        minutes = 0.0
        seconds = parts[0]

    else:
        raise ValueError(
            f"Invalid startTime: {startTime!r}"
        )

    if hours < 0 or minutes < 0 or seconds < 0:
        raise ValueError(
            f"startTime cannot be negative: {startTime!r}"
        )

    if minutes >= 60 or seconds >= 60:
        raise ValueError(
            "Minutes and seconds must be below 60: "
            f"{startTime!r}"
        )

    return (
        hours * 3600
        + minutes * 60
        + seconds
    )


def parseMovieTime(movieTime):
    """Convert an Excel time, number, or text value to seconds."""

    if hasattr(movieTime, "hour"):
        duration = (
            movieTime.hour * 3600
            + movieTime.minute * 60
            + movieTime.second
            + movieTime.microsecond / 1_000_000
        )

    else:
        try:
            duration = float(
                str(movieTime).strip()
            )

        except ValueError as error:
            raise ValueError(
                f"Invalid movieTime: {movieTime!r}"
            ) from error

    if duration <= 0:
        raise ValueError(
            "movieTime must be greater than zero: "
            f"{movieTime!r}"
        )

    return duration


def formatTimecode(seconds):
    """Convert seconds to HH:MM:SS.sss for FFmpeg."""

    hours = int(seconds // 3600)

    remainder = seconds - hours * 3600
    minutes = int(remainder // 60)
    secs = remainder - minutes * 60

    return (
        f"{hours:02d}:"
        f"{minutes:02d}:"
        f"{secs:06.3f}"
    )


def cleanMovieFile(movieFile):
    """Ensure the output file has a simple .mp4 filename."""

    movieFile = os.path.basename(
        str(movieFile).strip()
    )

    if not movieFile:
        raise ValueError(
            "movieFile cannot be blank"
        )

    if not movieFile.lower().endswith(".mp4"):
        movieFile += ".mp4"

    return movieFile


def extractAudio(videoPath, audioPath):
    """
    Extract the first audio stream from one MP4 file and save it as
    a 48 kHz, stereo, 16-bit PCM WAV file.
    """

    print(
        f"Extracting audio:\n"
        f"  Input:  {videoPath}\n"
        f"  Output: {audioPath}",
        flush=True
    )

    runCommand([
        "ffmpeg",
        "-y",
        "-i",
        str(videoPath),
        "-map",
        "0:a:0",
        "-vn",
        "-ar",
        audioSampleRate,
        "-ac",
        audioChannels,
        "-c:a",
        audioCodecPcm,
        str(audioPath),
    ])


# ----- Check dependencies and input files -----


checkBinary("ffmpeg")

if not spreadsheetPath.exists():
    raise FileNotFoundError(
        f"Spreadsheet not found: {spreadsheetPath}"
    )

if not sourceVideoPath.exists():
    raise FileNotFoundError(
        f"Source video not found: {sourceVideoPath}"
    )


# ----- Read the spreadsheet -----


workbook = load_workbook(
    spreadsheetPath,
    data_only=True
)

worksheet = workbook.active

headerRow = next(
    worksheet.iter_rows(
        min_row=1,
        max_row=1,
        values_only=True
    )
)

headers = [
    str(value).strip()
    if value is not None
    else ""
    for value in headerRow
]

requiredHeaders = {
    "startTime",
    "movieTime",
    "movieFile",
}

missingHeaders = requiredHeaders - set(headers)

if missingHeaders:
    raise ValueError(
        "Spreadsheet is missing required columns: "
        + ", ".join(sorted(missingHeaders))
    )

columnIndex = {
    header: index
    for index, header in enumerate(headers)
}


# ----- Cut clips and extract audio -----


successfulCount = 0
failedCount = 0

for rowNumber, values in enumerate(
    worksheet.iter_rows(
        min_row=2,
        values_only=True
    ),
    start=2
):

    # Skip fully blank spreadsheet rows
    if all(value is None for value in values):
        continue

    try:
        startTime = values[
            columnIndex["startTime"]
        ]

        movieTime = values[
            columnIndex["movieTime"]
        ]

        movieFile = values[
            columnIndex["movieFile"]
        ]

        startSeconds = parseTimeToSeconds(
            startTime
        )

        clipDuration = parseMovieTime(
            movieTime
        )

        movieFile = cleanMovieFile(
            movieFile
        )

        videoOutputPath = videoDir / movieFile

        audioFileName = videoOutputPath.stem + ".wav"
        audioOutputPath = audioDir / audioFileName

        print(
            f"\nRow {rowNumber}: {movieFile}\n"
            f"Start time: "
            f"{formatTimecode(startSeconds)}\n"
            f"Duration: {clipDuration} seconds\n"
            f"Video output: {videoOutputPath}\n"
            f"Audio output: {audioOutputPath}",
            flush=True
        )

        # ----- Cut the video clip -----

        runCommand([
            "ffmpeg",
            "-y",
            "-i",
            str(sourceVideoPath),
            "-ss",
            formatTimecode(startSeconds),
            "-t",
            str(clipDuration),
            "-c:v",
            videoCodec,
            "-crf",
            "18",
            "-preset",
            "medium",
            "-pix_fmt",
            pixelFormat,
            "-c:a",
            audioCodec,
            "-b:a",
            audioBitrate,
            "-movflags",
            "+faststart",
            str(videoOutputPath),
        ])

        # ----- Extract audio from the new clip -----

        extractAudio(
            videoOutputPath,
            audioOutputPath
        )

        successfulCount += 1

        print(
            f"Finished row {rowNumber}: "
            f"{movieFile}",
            flush=True
        )

    except Exception as error:
        failedCount += 1

        print(
            f"ERROR on row {rowNumber}: {error}",
            file=sys.stderr,
            flush=True
        )

        # Continue to the next spreadsheet row.
        continue


print("\nFinished processing spreadsheet.")
print(f"Successful rows: {successfulCount}")
print(f"Failed rows: {failedCount}")
print(f"Video output directory: {videoDir}")
print(f"Audio output directory: {audioDir}")